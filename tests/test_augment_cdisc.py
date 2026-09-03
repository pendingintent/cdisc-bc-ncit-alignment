"""Tests for src.augment_cdisc. Network calls are monkeypatched out."""

from __future__ import annotations

import pathlib
import sys

import pytest
from openpyxl import Workbook, load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src import augment_cdisc as ac  # noqa: E402

HEADERS = ["NCI Concept Code", "Preferred Name (name)"]


def _make_input(path: pathlib.Path, codes: list[str]) -> pathlib.Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(HEADERS)
    for code in codes:
        ws.append([code, f"Concept {code}"])
    wb.save(path)
    return path


def _fake_response() -> dict:
    return {
        "_links": {
            "biomedicalConcepts": [
                {"href": "/mdr/bc/biomedicalconcepts/C1", "title": "Glucose Measurement"},
            ]
        }
    }


def _patch_network(monkeypatch):
    monkeypatch.setattr(ac, "make_session", lambda: object())
    monkeypatch.setattr(ac, "fetch_biomedical_concepts", lambda session=None, timeout=120: _fake_response())


def test_augment_writes_new_columns_to_separate_output(tmp_path, monkeypatch):
    _patch_network(monkeypatch)
    input_path = _make_input(tmp_path / "in.xlsx", ["C1", "C2"])
    output_path = tmp_path / "out.xlsx"

    result = ac.augment(input_path, output_path)

    assert result == output_path
    wb = load_workbook(output_path, read_only=True)
    rows = list(wb["Sheet1"].iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS + ac.NEW_HEADERS)
    assert rows[1] == ("C1", "Concept C1", True, "/mdr/bc/biomedicalconcepts/C1", "Glucose Measurement")
    assert rows[2] == ("C2", "Concept C2", False, None, None)

    # input file must be untouched
    src_wb = load_workbook(input_path, read_only=True)
    src_rows = list(src_wb["Sheet1"].iter_rows(values_only=True))
    assert src_rows[0] == tuple(HEADERS)


def test_augment_in_place_replaces_input_atomically(tmp_path, monkeypatch):
    _patch_network(monkeypatch)
    input_path = _make_input(tmp_path / "in.xlsx", ["C1"])

    result = ac.augment(input_path)

    assert result == input_path
    wb = load_workbook(input_path, read_only=True)
    rows = list(wb["Sheet1"].iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS + ac.NEW_HEADERS)
    assert rows[1][2] is True

    # no leftover temp files
    leftovers = list(tmp_path.glob("*.xlsx.tmp"))
    assert leftovers == []


def test_augment_missing_input_raises(tmp_path, monkeypatch):
    _patch_network(monkeypatch)
    missing = tmp_path / "missing.xlsx"

    with pytest.raises(FileNotFoundError):
        ac.augment(missing, tmp_path / "out.xlsx")
