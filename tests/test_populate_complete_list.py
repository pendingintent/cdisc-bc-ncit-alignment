"""Tests for src.populate_complete_list. Network calls are monkeypatched out."""

from __future__ import annotations

import pathlib
import sys

from openpyxl import Workbook, load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src import populate_complete_list as pcl  # noqa: E402

HEADERS = [
    "NCI Concept Code",
    "Preferred Name (name)",
    "Concept  Status",
    "Semantic Type",
    "NCI Definition",
    "CDISC Definition",
    "Synonyms",
    "Parent  Concept Code",
    "Parent Concept Name",
    "Child Concept Code",
    "Child Concept Name",
    "Active (True/False)",
]


def _make_template(path: pathlib.Path) -> pathlib.Path:
    wb = Workbook()
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active
    ws.title = "Sheet1"
    ws.append(HEADERS)
    wb.save(path)
    return path


def _concept(code: str) -> dict:
    return {"code": code, "name": f"Concept {code}", "conceptStatus": "DEFAULT", "active": True}


def _patch_network(monkeypatch, codes: list[str]):
    monkeypatch.setattr(pcl, "make_session", lambda pool=8: object())
    monkeypatch.setattr(pcl, "fetch_all_codes", lambda session=None, timeout=120: list(codes))

    def fake_fetch_concepts(codes, session=None, batch_size=200, max_workers=8, on_progress=None):
        total = (len(codes) + batch_size - 1) // batch_size or 1
        for i, code in enumerate(codes, start=1):
            if on_progress and i % batch_size == 0:
                on_progress(i // batch_size, total)
            yield _concept(code)

    monkeypatch.setattr(pcl, "fetch_concepts", fake_fetch_concepts)


def test_populate_writes_headers_and_rows(tmp_path, monkeypatch):
    template = _make_template(tmp_path / "template.xlsx")
    output = tmp_path / "out.xlsx"
    _patch_network(monkeypatch, ["C1", "C2", "C3"])

    result = pcl.populate(output_path=output, template_path=template, batch_size=2, workers=2)

    assert result == output
    assert output.exists()
    wb = load_workbook(output, read_only=True)
    rows = list(wb["Sheet1"].iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS)
    assert [r[0] for r in rows[1:]] == ["C1", "C2", "C3"]
    assert rows[1][1] == "Concept C1"


def test_populate_respects_limit(tmp_path, monkeypatch):
    template = _make_template(tmp_path / "template.xlsx")
    output = tmp_path / "out.xlsx"
    _patch_network(monkeypatch, ["C1", "C2", "C3", "C4", "C5"])

    pcl.populate(output_path=output, template_path=template, limit=2)

    wb = load_workbook(output, read_only=True)
    rows = list(wb["Sheet1"].iter_rows(values_only=True))
    assert [r[0] for r in rows[1:]] == ["C1", "C2"]


def test_populate_creates_missing_output_dir(tmp_path, monkeypatch):
    template = _make_template(tmp_path / "template.xlsx")
    output = tmp_path / "nested" / "dir" / "out.xlsx"
    _patch_network(monkeypatch, ["C1"])

    pcl.populate(output_path=output, template_path=template)

    assert output.exists()
